import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(student_name="Anas Mehfooz", matr_no="[Your Matriculation Number]"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    NAVY = RGBColor(26, 43, 76)
    BLUE = RGBColor(41, 128, 185)
    TEAL = RGBColor(22, 160, 133)
    DARK_GRAY = RGBColor(44, 62, 80)
    LIGHT_BG = RGBColor(245, 247, 250)
    WHITE = RGBColor(255, 255, 255)
    BORDER_GRAY = RGBColor(218, 225, 231)
    
    blank_slide_layout = prs.slide_layouts[6]
    
    def add_header(slide, title_text, category_text="DLBCSPJWD01 - PORTFOLIO PHASE 1"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = TEAL
        
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = BLUE
        line.line.color.rgb = BLUE

    def add_card(slide, left, top, width, height, title, content_list, bg_color=LIGHT_BG, border_color=BORDER_GRAY):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.25)
        
        p_head = tf.paragraphs[0]
        p_head.text = title
        p_head.font.size = Pt(16)
        p_head.font.bold = True
        p_head.font.color.rgb = NAVY
        p_head.space_after = Pt(12)
        
        for item in content_list:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(13)
            p.font.color.rgb = DARK_GRAY
            p.space_after = Pt(6)

    # TITLE SLIDE
    slide_title = prs.slides.add_slide(blank_slide_layout)
    bg = slide_title.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    
    accent_box = slide_title.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
    accent_box.fill.solid()
    accent_box.fill.fore_color.rgb = WHITE
    accent_box.line.color.rgb = TEAL
    accent_box.line.width = Pt(2)
    
    tf_t = accent_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = Inches(0.6)
    tf_t.margin_top = Inches(0.6)
    
    p0 = tf_t.paragraphs[0]
    p0.text = "IU INTERNATIONALE HOCHSCHULE"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = TEAL
    p0.space_after = Pt(20)
    
    p1 = tf_t.add_paragraph()
    p1.text = "WeatherSync - Smart Activity & Task Planner"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    p1.space_after = Pt(10)
    
    p2 = tf_t.add_paragraph()
    p2.text = "Portfolio Phase 1: Conceptual Design Presentation"
    p2.font.size = Pt(20)
    p2.font.color.rgb = BLUE
    p2.space_after = Pt(30)
    
    p3 = tf_t.add_paragraph()
    p3.text = f"Course: Project Java and Web Development (DLBCSPJWD01)\nStudent Name: {student_name}\nMatriculation Number: {matr_no}"
    p3.font.size = Pt(14)
    p3.font.color.rgb = DARK_GRAY

    # SLIDE 1
    slide1 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide1, "Slide 1: Envisioned Web Application Concept & Purpose")
    add_card(slide1, 0.8, 1.8, 5.6, 5.0, "Core Application Idea", [
        "• Name: WeatherSync - Weather-Aware Task & Activity Planner.",
        "• Core Purpose: Web app connecting daily task management with real-time weather forecasts.",
        "• Problem Addressed: Users misplan outdoor tasks when checking weather apps separately from planners.",
        "• Solution: Automatically checks weather metrics and calculates activity suitability scores."
    ])
    add_card(slide1, 6.9, 1.8, 5.6, 5.0, "Key Features & Vision", [
        "• Weather Dashboard: Real-time 5-day weather view synced with task schedule.",
        "• Suitability Engine: Backend algorithm scoring outdoor tasks against forecast data.",
        "• Location Search: Weather lookup across multiple cities.",
        "• Responsive Web UI: Smooth layout on smartphones, tablets, and desktops."
    ])

    # SLIDE 2
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "Slide 2: Target User Groups & Primary Use Cases")
    add_card(slide2, 0.8, 1.8, 3.6, 5.0, "1. Active Individuals", [
        "• Profile: Runners, cyclists, outdoor fitness enthusiasts.",
        "• Needs: Hourly forecasts, rain warnings, dry weather windows.",
        "• Use Case: Plan a 10km run during the dryest hour of the afternoon."
    ])
    add_card(slide2, 4.8, 1.8, 3.6, 5.0, "2. Remote Workers", [
        "• Profile: People balancing work, outdoor breaks, and chores.",
        "• Needs: Efficient daily scheduling combining work and outdoor breaks.",
        "• Use Case: Schedule lawn mowing or errands during sunny breaks."
    ])
    add_card(slide2, 8.8, 1.8, 3.6, 5.0, "3. Event Planners", [
        "• Profile: Families and groups organizing weekend trips or outdoor events.",
        "• Needs: Multi-city weather lookup and flexible scheduling.",
        "• Use Case: Compare weather across 3 nearby destinations."
    ])

    # SLIDE 3
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "Slide 3: Key User Benefits & Value Proposition")
    add_card(slide3, 0.8, 1.8, 5.6, 2.3, "Time Savings", [
        "• Eliminates switching between separate calendar, task manager, and weather apps.",
        "• Instant weather-suitability tags on every planned activity."
    ])
    add_card(slide3, 6.9, 1.8, 5.6, 2.3, "Smart Recommendations", [
        "• Backend logic warns users of rain or poor weather before outdoor tasks.",
        "• Recommends alternative indoor time slots."
    ])
    add_card(slide3, 0.8, 4.4, 5.6, 2.4, "Multi-Device Compatibility", [
        "• Fully responsive layout designed with modern CSS (Flexbox & Grid).",
        "• Provides clear readability on mobile, tablet, and desktop displays."
    ])
    add_card(slide3, 6.9, 4.4, 5.6, 2.4, "Clean Interface", [
        "• Clear visual layout with weather icons, temperature badges, and task cards.",
        "• Quick search and filter features for tasks."
    ])

    # SLIDE 4
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "Slide 4: Technical Architecture & Technology Stack")
    add_card(slide4, 0.8, 1.8, 3.6, 5.0, "Frontend Layer", [
        "• HTML5: Semantic markup structure.",
        "• CSS3: Custom CSS variables, Flexbox/Grid layout, responsive media queries.",
        "• JavaScript (ES6+): Dynamic DOM updates, async fetch requests."
    ])
    add_card(slide4, 4.8, 1.8, 3.6, 5.0, "Backend Layer", [
        "• Node.js: Asynchronous JavaScript runtime.",
        "• Express.js: RESTful API server routing.",
        "• Storage: Local JSON file storage (`tasks.json`)."
    ])
    add_card(slide4, 8.8, 1.8, 3.6, 5.0, "External API Integration", [
        "• External API: Open-Meteo REST API.",
        "• Backend Proxy: Weather API is called strictly from the Node.js server.",
        "• Server transforms raw weather data into activity suitability scores for the frontend."
    ])

    # SLIDE 5
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "Slide 5: Functional Scope & Dynamic Features")
    add_card(slide5, 0.8, 1.8, 5.6, 5.0, "Dynamic Interaction 1: Weather Lookup", [
        "1. User enters location/city name.",
        "2. Client sends HTTP GET request to `/api/weather?city=X`.",
        "3. Backend calls Open-Meteo API and calculates suitability score.",
        "4. Backend returns formatted JSON payload.",
        "5. Client JS dynamically renders forecast widgets without page reload."
    ])
    add_card(slide5, 6.9, 1.8, 5.6, 5.0, "Dynamic Interaction 2: Task CRUD Operations", [
        "1. User creates, updates, or completes a task.",
        "2. Client JS posts HTTP request to `/api/tasks`.",
        "3. Backend validates input and saves to local storage.",
        "4. Backend returns updated task array.",
        "5. Client dynamically updates task list UI instantly."
    ])

    output_path = os.path.join("d:\\My_project", "Anas-Mehfooz_DLBCSPJWD01_P1_S.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()
